from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".html", ".css", ".sql", ".yaml", ".yml"}
IMAGE_EXT = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}
MAX_PAGES = 1000
MAX_CHARS = 5_000_000


class FileRejected(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


def classify(filename: str) -> str | None:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXT:
        return "image"
    if ext in TEXT_EXT or ext in {".pdf", ".docx"}:
        return "document"
    return None


def sniff_image(data: bytes, ext: str) -> bool:
    if ext == ".png":
        return data.startswith(b"\x89PNG\r\n\x1a\n")
    if ext in {".jpg", ".jpeg"}:
        return data.startswith(b"\xff\xd8\xff")
    if ext == ".gif":
        return data.startswith((b"GIF87a", b"GIF89a"))
    if ext == ".webp":
        return data[:4] == b"RIFF" and data[8:12] == b"WEBP"
    return False


def parse_document(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Return (page, text) sections. Raises FileRejected for unreadable/unsupported files."""
    ext = Path(filename).suffix.lower()
    sections: list[tuple[int | None, str]]

    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise FileRejected(422, "This file is not a valid PDF.")
        try:
            reader = PdfReader(BytesIO(data))
            sections = [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages[:MAX_PAGES])]
        except Exception as exc:  # malformed / encrypted
            raise FileRejected(422, "Could not read this PDF (it may be encrypted or damaged).") from exc
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise FileRejected(422, "This file is not a valid DOCX.")
        try:
            doc = DocxDocument(BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            sections = [(None, "\n\n".join(parts))]
        except Exception as exc:
            raise FileRejected(422, "Could not read this DOCX file.") from exc
    elif ext in TEXT_EXT:
        sections = [(None, data.decode("utf-8", errors="replace"))]
    else:
        raise FileRejected(415, "Unsupported file type. Use PDF, DOCX, TXT, MD, CSV or JSON (or an image in chat).")

    total = 0
    cleaned: list[tuple[int | None, str]] = []
    for page, text in sections:
        text = text.replace("\x00", "").strip()
        if text:
            total += len(text)
            if total > MAX_CHARS:
                break
            cleaned.append((page, text))
    if not cleaned:
        raise FileRejected(422, "No readable text found. Scanned PDFs need OCR, which is planned for a later phase.")
    return cleaned
